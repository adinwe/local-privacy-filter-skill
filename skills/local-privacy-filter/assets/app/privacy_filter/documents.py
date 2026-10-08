"""Bounded, offline document-to-clean-document conversion.

Only selected text and basic structure are reconstructed. Original ZIP parts,
PDF objects, properties, links, images and other embedded content never enter
the output. The caller must review the resulting text before sharing it.
"""

from __future__ import annotations

import csv
import io
import logging
import math
import os
import re
import shutil
import subprocess
import time
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time
from pathlib import Path, PurePosixPath
from typing import Callable, Iterable, Iterator, Sequence

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_CHARS = 100_000
SUPPORTED_EXTENSIONS = frozenset({"csv", "xlsx", "docx", "pdf"})
MAX_ZIP_MEMBERS = 2_000
MAX_ZIP_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 16 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_XML_ELEMENTS = 200_000
MAX_ROWS = 10_000
MAX_COLUMNS = 200
MAX_CELLS = 30_000
MAX_SHEETS = 20
MAX_PAGES = 50
MAX_OCR_PAGES = 20
MAX_PDF_STREAM_BYTES = 8 * 1024 * 1024
MAX_OCR_IMAGE_BYTES = 20 * 1024 * 1024
OCR_DOCUMENT_SECONDS = 120

_UNSUPPORTED = "Choose a CSV, XLSX, DOCX or PDF file. Legacy and macro-enabled files are unsupported."
_EMPTY = "The uploaded file is empty."
_TOO_LARGE = "The file exceeds the 20 MiB upload limit."
_TEXT_LIMIT = "The document exceeds the 100,000-character processing limit."
_STRUCTURE_LIMIT = "The document has too many pages, rows, cells or embedded parts to process safely."
_ZIP_LIMIT = "The document's compressed contents exceed safe processing limits."
_MALFORMED = "The document could not be read. It may be damaged or unsupported."
_ENCRYPTED = "Encrypted or password-protected documents are unsupported."
_MACROS = "Macro-enabled documents are unsupported. Save a macro-free copy first."
_ENCODING = "The CSV must use UTF-8 or a UTF-16 byte-order mark."
_FILTER_FAILED = "Local filtering failed; no downloadable file was produced."
_OCR_MISSING = "This PDF needs local OCR, which is unavailable. No file was produced."
_OCR_FAILED = "Local PDF OCR could not read a page safely. No file was produced."
_OCR_LIMIT = "The PDF requires too many OCR pages or too much processing time."
SAFE_ERROR_MESSAGES = frozenset({
    _UNSUPPORTED, _EMPTY, _TOO_LARGE, _TEXT_LIMIT, _STRUCTURE_LIMIT,
    _ZIP_LIMIT, _MALFORMED, _ENCRYPTED, _MACROS, _ENCODING,
    _FILTER_FAILED, _OCR_MISSING, _OCR_FAILED, _OCR_LIMIT,
})

_REVIEW = "Review the filtered file before sharing: automated detection may miss identifiers or mask harmless text."
_CSV_WARNING = "CSV text and cell values were rebuilt. Formatting and the original encoding were simplified; spreadsheet formulas were neutralized."
_XLSX_WARNING = "A new workbook contains filtered cell text from all sheets, including hidden sheets. Sheet names, comments, links, images, charts, metadata and original formatting were omitted; formulas were omitted and typed values may be text."
_DOCX_WARNING = "A new Word document contains filtered body paragraphs and tables. Headers, footers, comments, deleted revisions, fields, links, images, text boxes, embedded objects, metadata and original formatting were omitted; inserted revision text was flattened."
_PDF_WARNING = "A new text PDF contains filtered extracted page text. Original layout, images, annotations, forms, links, attachments and metadata were omitted; non-text content was not preserved."
_OCR_WARNING = "Image-containing or non-readable PDF pages were processed using local English OCR. Mixed pages may repeat readable text; OCR may misread identifiers. Review every page before sharing."


class DocumentError(Exception):
    """Only fixed safe messages may be exposed to the browser."""

    def __init__(self, message: str, status: int = 400) -> None:
        safe_message = message if message in SAFE_ERROR_MESSAGES else _MALFORMED
        super().__init__(safe_message)
        self.status = status
        self.status_code = status


@dataclass(frozen=True)
class DocumentResult:
    content: bytes
    media_type: str
    extension: str
    entity_count: int
    warnings: list[str]
    units: int


@dataclass
class _Budget:
    characters: int = 0
    cells: int = 0
    rows: int = 0

    def text(self, amount: int) -> None:
        self.characters += amount
        if self.characters > MAX_DOCUMENT_CHARS:
            raise DocumentError(_TEXT_LIMIT, 413)

    def row(self, cells: int) -> None:
        self.rows += 1
        self.cells += cells
        if self.rows > MAX_ROWS or self.cells > MAX_CELLS or cells > MAX_COLUMNS:
            raise DocumentError(_STRUCTURE_LIMIT, 413)


def _silence_parser_logs(package: str) -> None:
    # Malformed PDF font names can contain user values in parser warning logs.
    # Silence only the parser package, preserving the app's operational logging.
    parent = logging.getLogger(package)
    parent.setLevel(logging.CRITICAL + 1)
    parent.propagate = False
    if not parent.handlers:
        parent.addHandler(logging.NullHandler())
    for name in list(logging.Logger.manager.loggerDict):
        if name == package or name.startswith(package + "."):
            logger = logging.getLogger(name)
            logger.disabled = True


def _mask_text(text: str, findings: list[dict]) -> tuple[str, int]:
    """Union spans before replacement so overlapping labels cannot leak tails."""
    spans = sorted(findings, key=lambda item: (item["start"], item["end"]))
    merged: list[dict] = []
    for finding in spans:
        item = {"start": finding["start"], "end": finding["end"], "entity_type": finding["entity_type"]}
        if merged and item["start"] <= merged[-1]["end"]:
            previous = merged[-1]
            previous["end"] = max(previous["end"], item["end"])
            if previous["entity_type"] != item["entity_type"]:
                previous["entity_type"] = "SENSITIVE_DATA"
        else:
            merged.append(item)
    result: list[str] = []
    position = 0
    for item in merged:
        result.extend((text[position:item["start"]], f'[{item["entity_type"]}]'))
        position = item["end"]
    result.append(text[position:])
    return "".join(result), len(merged)


def _filter_values(
    values: Sequence[str], redact: Callable[[str], dict], budget: _Budget,
    headers: Sequence[str] | None = None,
) -> tuple[list[str], int]:
    """Keep row/paragraph context, but apply only spans inside retained values.

    A header may be detected as part of context again; those context-only spans
    are never copied into a value. Findings must use the original offsets.
    """
    pieces: list[str] = []
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for index, value in enumerate(values):
        if index:
            pieces.append("\n")
            cursor += 1
        if headers is not None and index < len(headers) and headers[index]:
            context = headers[index] + ": "
            pieces.append(context)
            cursor += len(context)
        offsets.append((cursor, cursor + len(value)))
        pieces.append(value)
        cursor += len(value)
    text = "".join(pieces)
    budget.text(len(text))
    if not text.strip():
        return list(values), 0
    try:
        response = redact(text)
        if not isinstance(response, dict) or not isinstance(response.get("findings"), list):
            raise ValueError("missing spans")
        findings = response["findings"]
        for item in findings:
            if (
                not isinstance(item, dict)
                or type(item.get("start")) is not int
                or type(item.get("end")) is not int
                or not 0 <= item["start"] < item["end"] <= len(text)
                or not isinstance(item.get("entity_type"), str)
                or re.fullmatch(r"[A-Z][A-Z0-9_]{0,63}", item["entity_type"]) is None
            ):
                raise ValueError("invalid spans")
    except Exception:
        raise DocumentError(_FILTER_FAILED, 503) from None
    outputs: list[str] = []
    count = 0
    for value, (start, end) in zip(values, offsets):
        local = [
            {"start": max(item["start"], start) - start,
             "end": min(item["end"], end) - start,
             "entity_type": item["entity_type"]}
            for item in findings if item["start"] < end and item["end"] > start
        ]
        masked, detected = _mask_text(value, local)
        outputs.append(masked)
        count += detected
    return outputs, count


def _filter_rows(
    rows: Iterable[Sequence[str]], redact: Callable[[str], dict], budget: _Budget,
) -> Iterator[tuple[list[str], int]]:
    """Batch rows with column context; avoid a model call for each cell or row."""
    pending: list[list[str]] = []
    contexts: list[str] = []
    headers: list[str] | None = None
    characters = 0

    def flush() -> list[tuple[list[str], int]]:
        flat = [value for row in pending for value in row]
        masked, count = _filter_values(flat, redact, budget, contexts)
        result: list[tuple[list[str], int]] = []
        position = 0
        for index, row in enumerate(pending):
            result.append((masked[position:position + len(row)], count if index == 0 else 0))
            position += len(row)
        return result

    for source_row in rows:
        row = list(source_row)
        budget.row(len(row))
        context = [headers[i] if headers is not None and i < len(headers) else "" for i in range(len(row))]
        size = sum(len(value) + (len(label) + 2 if label else 0) + 1 for value, label in zip(row, context))
        if size > MAX_DOCUMENT_CHARS:
            raise DocumentError(_TEXT_LIMIT, 413)
        if pending and characters + size > 12_000:
            yield from flush()
            pending.clear()
            contexts.clear()
            characters = 0
        pending.append(row)
        contexts.extend(context)
        characters += size
        if headers is None and any(row):
            headers = row
    if pending:
        yield from flush()


def _zip_preflight(data: bytes, extension: str) -> None:
    """Bound ZIP inflation and validate XML without resolving DTD/entities."""
    from defusedxml import ElementTree as safe_xml

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_MEMBERS:
            raise DocumentError(_STRUCTURE_LIMIT, 413)
        names: set[str] = set()
        total = 0
        for item in entries:
            name = item.filename
            path = PurePosixPath(name)
            if name in names or path.is_absolute() or ".." in path.parts or "\\" in name or "\x00" in name:
                raise DocumentError(_MALFORMED)
            names.add(name)
            if item.flag_bits & 1:
                raise DocumentError(_ENCRYPTED, 422)
            if "vbaproject" in name.lower() or name.lower().endswith((".bin", ".vba")):
                # Embedded binary objects are omitted, but a VBA payload must be rejected.
                if "vba" in name.lower():
                    raise DocumentError(_MACROS, 422)
            total += item.file_size
            if (
                item.file_size > MAX_ZIP_MEMBER_BYTES
                or total > MAX_ZIP_UNCOMPRESSED_BYTES
                or item.file_size > max(1, item.compress_size) * MAX_ZIP_RATIO
            ):
                raise DocumentError(_ZIP_LIMIT, 413)
            if name.endswith((".xml", ".rels")):
                payload = archive.read(item)
                count = depth = 0
                for event, element in safe_xml.iterparse(io.BytesIO(payload), events=("start", "end"), forbid_dtd=True):
                    if event == "start":
                        count += 1
                        depth += 1
                        if count > MAX_XML_ELEMENTS or depth > 128:
                            raise DocumentError(_STRUCTURE_LIMIT, 413)
                        if extension == "xlsx" and name.startswith("xl/worksheets/"):
                            tag = element.tag.rsplit("}", 1)[-1]
                            coordinate = element.attrib.get("r", "")
                            if tag == "row" and coordinate and (not coordinate.isdecimal() or not 1 <= int(coordinate) <= MAX_ROWS):
                                raise DocumentError(_STRUCTURE_LIMIT, 413)
                            if tag == "c" and coordinate:
                                match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", coordinate)
                                if match is None:
                                    raise DocumentError(_MALFORMED)
                                column = 0
                                for character in match[1]:
                                    column = column * 26 + ord(character) - ord("A") + 1
                                if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS:
                                    raise DocumentError(_STRUCTURE_LIMIT, 413)
                    else:
                        depth -= 1
                        element.clear()
                if name == "[Content_Types].xml" and b"macroenabled" in payload.lower():
                    raise DocumentError(_MACROS, 422)
        required = "word/document.xml" if extension == "docx" else "xl/workbook.xml"
        if required not in names or "[Content_Types].xml" not in names:
            raise DocumentError(_MALFORMED)


def _csv_formula_safe(value: str) -> str:
    # Quotes alone do not neutralize spreadsheet formulas upon CSV opening.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _csv(data: bytes, redact: Callable[[str], dict]) -> DocumentResult:
    try:
        encoding = "utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = data.decode(encoding)
    except UnicodeError:
        raise DocumentError(_ENCODING, 422) from None
    if "\x00" in text:
        raise DocumentError(_MALFORMED)
    if len(text) > MAX_DOCUMENT_CHARS:
        raise DocumentError(_TEXT_LIMIT, 413)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    budget = _Budget()
    count = units = 0
    for filtered, found in _filter_rows(rows, redact, budget):
        writer.writerow([_csv_formula_safe(value) for value in filtered])
        count += found
        units += 1
    return DocumentResult(output.getvalue().encode("utf-8"), "text/csv; charset=utf-8", "csv", count, [_REVIEW, _CSV_WARNING], units)


def _cell_string(cell: object) -> str:
    value = cell.value
    if cell.data_type == "f":
        return "[FORMULA OMITTED]"
    if value is None:
        return ""
    if isinstance(value, datetime):
        from openpyxl.styles.numbers import is_datetime
        return value.date().isoformat() if is_datetime(cell.number_format) == "date" else value.isoformat()
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise DocumentError(_MALFORMED)
        return str(value)
    if not isinstance(value, str):
        raise DocumentError(_MALFORMED)
    return value


def _xlsx(data: bytes, redact: Callable[[str], dict]) -> DocumentResult:
    from openpyxl import Workbook, load_workbook
    from openpyxl.packaging.core import DocumentProperties

    _silence_parser_logs("openpyxl")
    original = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(original.worksheets) > MAX_SHEETS:
            raise DocumentError(_STRUCTURE_LIMIT, 413)
        fresh = Workbook()
        fresh.remove(fresh.active)
        fresh.properties = DocumentProperties(creator="", lastModifiedBy="", title="", subject="", description="", keywords="", category="", created=datetime(2000, 1, 1), modified=datetime(2000, 1, 1))
        budget = _Budget()
        count = units = 0
        for sheet_number, sheet in enumerate(original.worksheets, 1):
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise DocumentError(_STRUCTURE_LIMIT, 413)
            target = fresh.create_sheet(f"Sheet{sheet_number}")
            # Do not let an incorrect XML dimension silently hide cells.
            sheet.reset_dimensions()
            values = ([_cell_string(cell) for cell in cells] for cells in sheet.iter_rows())
            for row_number, (filtered, found) in enumerate(_filter_rows(values, redact, budget), 1):
                for column, value in enumerate(filtered, 1):
                    target_cell = target.cell(row=row_number, column=column, value=value)
                    # Never let a string beginning with '=' become a formula.
                    target_cell.data_type = "s"
                units += 1
                count += found
        if not fresh.worksheets:
            fresh.create_sheet("Sheet1")
        output = io.BytesIO()
        fresh.save(output)
    finally:
        original.close()
    return DocumentResult(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx", count, [_REVIEW, _XLSX_WARNING], units)


def _docx_paragraph_text(element: object) -> str:
    from docx.oxml.ns import qn

    omitted = {qn("w:del"), qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:txbxContent")}
    pieces: list[str] = []

    def walk(node: object) -> None:
        if node.tag in omitted:
            return
        if node.tag == qn("w:t"):
            pieces.append(node.text or "")
        elif node.tag == qn("w:tab"):
            pieces.append("\t")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            pieces.append("\n")
        else:
            for child in node:
                walk(child)

    walk(element)
    return "".join(pieces)


def _docx(data: bytes, redact: Callable[[str], dict]) -> DocumentResult:
    from docx import Document
    from docx.oxml.ns import qn

    _silence_parser_logs("docx")
    original = Document(io.BytesIO(data))
    fresh = Document()
    props = fresh.core_properties
    for name in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category", "content_status", "identifier", "language", "version"):
        setattr(props, name, "")
    props.created = props.modified = datetime(2000, 1, 1)
    props.revision = 1
    budget = _Budget()
    count = units = 0
    pending: list[str] = []

    def flush_paragraphs() -> None:
        nonlocal count, units
        if not pending:
            return
        filtered, found = _filter_values(pending, redact, budget)
        for text in filtered:
            fresh.add_paragraph(text)
        count += found
        units += len(pending)
        pending.clear()

    for node in original.element.body:
        if node.tag == qn("w:p"):
            pending.append(_docx_paragraph_text(node))
            if len(pending) > MAX_ROWS:
                raise DocumentError(_STRUCTURE_LIMIT, 413)
            if sum(map(len, pending)) + len(pending) - 1 > MAX_DOCUMENT_CHARS:
                raise DocumentError(_TEXT_LIMIT, 413)
        elif node.tag == qn("w:tbl"):
            flush_paragraphs()
            rows: list[list[str]] = []
            for row in node.findall(qn("w:tr")):
                cells: list[str] = []
                for cell in row.findall(qn("w:tc")):
                    # Nested table text is flattened into this cell, then filtered.
                    paragraphs = []
                    for paragraph in cell.iter(qn("w:p")):
                        ancestor = paragraph.getparent()
                        omit = False
                        while ancestor is not None and ancestor is not cell:
                            if ancestor.tag in {qn("w:del"), qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:txbxContent")}:
                                omit = True
                                break
                            ancestor = ancestor.getparent()
                        if not omit:
                            paragraphs.append(paragraph)
                    cells.append("\n".join(_docx_paragraph_text(p) for p in paragraphs))
                rows.append(cells)
                if len(rows) > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise DocumentError(_STRUCTURE_LIMIT, 413)
            if not rows:
                continue
            width = max(map(len, rows))
            target = fresh.add_table(rows=len(rows), cols=width)
            for row_number, (filtered, found) in enumerate(_filter_rows(rows, redact, budget)):
                count += found
                units += len(filtered)
                for column, value in enumerate(filtered):
                    target.cell(row_number, column).text = value
    flush_paragraphs()
    output = io.BytesIO()
    fresh.save(output)
    return DocumentResult(output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx", count, [_REVIEW, _DOCX_WARNING], units)


def _ocr_binaries() -> tuple[str, str] | None:
    result: list[str] = []
    for name in ("pdftoppm", "tesseract"):
        candidate = Path("/opt/homebrew/bin") / name
        path = str(candidate) if candidate.is_file() and os.access(candidate, os.X_OK) else shutil.which(name)
        if not path:
            return None
        result.append(path)
    return result[0], result[1]


def ocr_available() -> bool:
    return _ocr_binaries() is not None


def _ocr_page(data: bytes, page_number: int, deadline: float) -> str:
    binaries = _ocr_binaries()
    if binaries is None:
        raise DocumentError(_OCR_MISSING, 422)
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise DocumentError(_OCR_LIMIT, 413)
    try:
        rendered = subprocess.run(
            [binaries[0], "-f", str(page_number), "-l", str(page_number), "-scale-to", "2000", "-singlefile", "-png", "-"],
            input=data, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            timeout=min(30, remaining), check=True,
        ).stdout
        if not rendered.startswith(b"\x89PNG\r\n\x1a\n") or len(rendered) > MAX_OCR_IMAGE_BYTES:
            raise DocumentError(_OCR_FAILED, 422)
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise DocumentError(_OCR_LIMIT, 413)
        recognized = subprocess.run(
            [binaries[1], "stdin", "stdout", "-l", "eng", "--psm", "3"],
            input=rendered, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            timeout=min(30, remaining), check=True,
        ).stdout
        if len(recognized) > MAX_DOCUMENT_CHARS * 4:
            raise DocumentError(_TEXT_LIMIT, 413)
        text = recognized.decode("utf-8").strip()
        if not text:
            raise DocumentError(_OCR_FAILED, 422)
        return text
    except DocumentError:
        raise
    except (subprocess.TimeoutExpired, subprocess.CalledProcessError, OSError, UnicodeError):
        raise DocumentError(_OCR_FAILED, 422) from None


def _pdf_page_content(page: object) -> tuple[bool, bool]:
    """Find image XObjects (including nested forms/inline images) without decoding pixels."""
    from pypdf.generic import ContentStream

    seen: set[int] = set()
    total = 0
    has_image = False
    has_content = False

    def visit(stream: object, resources: object, depth: int = 0) -> None:
        nonlocal total, has_image, has_content
        if depth > 20:
            raise DocumentError(_STRUCTURE_LIMIT, 413)
        if stream is not None:
            decoded = stream.get_data()
            total += len(decoded)
            if total > MAX_PDF_STREAM_BYTES:
                raise DocumentError(_STRUCTURE_LIMIT, 413)
            content = stream if isinstance(stream, ContentStream) else ContentStream(stream, page.pdf)
            operations = content.operations
            # Empty pages often still contain font/transform setup instructions.
            painting = {b"Tj", b"TJ", b"'", b'"', b"Do", b"INLINE IMAGE", b"S", b"s", b"f", b"F", b"f*", b"B", b"B*", b"b", b"b*", b"sh"}
            has_content = has_content or any(operator in painting for _, operator in operations)
            if any(operator == b"INLINE IMAGE" for _, operator in operations):
                has_image = True
        if not resources:
            return
        resources = resources.get_object()
        xobjects = resources.get("/XObject")
        if not xobjects:
            return
        xobjects = xobjects.get_object()
        if len(xobjects) > 1_000:
            raise DocumentError(_STRUCTURE_LIMIT, 413)
        for reference in xobjects.values():
            obj = reference.get_object()
            identity = id(obj)
            if identity in seen:
                continue
            seen.add(identity)
            if len(seen) > 1_000:
                raise DocumentError(_STRUCTURE_LIMIT, 413)
            if obj.get("/Subtype") == "/Image":
                has_image = True
            elif obj.get("/Subtype") == "/Form":
                visit(obj, obj.get("/Resources", resources), depth + 1)

    visit(page.get_contents(), page.get("/Resources"))
    return has_image, has_content


def _build_text_pdf(pages: Sequence[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    import reportlab

    font_name = "PrivacyFilterVera"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        font_file = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
        pdfmetrics.registerFont(TTFont(font_name, str(font_file)))
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4, invariant=1, pageCompression=1)
    pdf.setAuthor("")
    pdf.setTitle("Filtered document")
    pdf.setSubject("")
    pdf.setCreator("Local Privacy Filter")
    width, height = A4
    margin = 45
    for page_index, page in enumerate(pages):
        if page_index:
            pdf.showPage()
        pdf.setFont(font_name, 10)
        y = height - margin
        for paragraph in page.splitlines() or [""]:
            # Discard unprintable extraction controls; they carry no rendered text.
            paragraph = "".join(char for char in paragraph if char >= " " or char == "\t")
            lines = simpleSplit(paragraph, font_name, 10, width - 2 * margin) or [""]
            for line in lines:
                if y < margin:
                    pdf.showPage()
                    pdf.setFont(font_name, 10)
                    y = height - margin
                pdf.drawString(margin, y, line)
                y -= 14
    pdf.save()
    return output.getvalue()


def _pdf(data: bytes, redact: Callable[[str], dict]) -> DocumentResult:
    from pypdf import PdfReader, apply_configuration

    _silence_parser_logs("pypdf")
    pages: list[str] = []
    count = ocr_pages = 0
    budget = _Budget()
    deadline = time.monotonic() + OCR_DOCUMENT_SECONDS
    # Modern pypdf configuration is per-context, avoiding global parser mutations.
    with apply_configuration(
        maximum_declared_stream_length=MAX_PDF_STREAM_BYTES,
        array_based_stream_maximum_output_length=MAX_PDF_STREAM_BYTES,
        zlib_maximum_output_length=MAX_PDF_STREAM_BYTES,
        lzw_maximum_output_length=MAX_PDF_STREAM_BYTES,
        run_length_maximum_output_length=MAX_PDF_STREAM_BYTES,
        jbig2_maximum_output_length=MAX_OCR_IMAGE_BYTES,
        image_maximum_buffer_size=MAX_OCR_IMAGE_BYTES,
        page_tree_maximum_entries=MAX_PAGES + 1,
        xform_maximum_invocations_per_extraction=1_000,
        disable_legacy_handling=True,
    ):
        reader = PdfReader(io.BytesIO(data), strict=True, root_object_recovery_limit=1_000)
        if reader.is_encrypted:
            raise DocumentError(_ENCRYPTED, 422)
        if not 0 < len(reader.pages) <= MAX_PAGES:
            raise DocumentError(_STRUCTURE_LIMIT, 413)
        for number, page in enumerate(reader.pages, 1):
            has_image, has_content = _pdf_page_content(page)
            text = page.extract_text() or ""
            if len(text) > MAX_DOCUMENT_CHARS:
                raise DocumentError(_TEXT_LIMIT, 413)
            if has_image or (has_content and not text.strip()):
                ocr_pages += 1
                if ocr_pages > MAX_OCR_PAGES:
                    raise DocumentError(_OCR_LIMIT, 413)
                recognized = _ocr_page(data, number, deadline)
                # Keep readable text as well: OCR must not erase an existing identifier.
                text = text + "\n" + recognized if text.strip() and text.strip() != recognized.strip() else recognized
            filtered, found = _filter_values([text], redact, budget)
            pages.append(filtered[0])
            count += found
    warnings = [_REVIEW, _PDF_WARNING]
    if ocr_pages:
        warnings.append(_OCR_WARNING)
    return DocumentResult(_build_text_pdf(pages), "application/pdf", "pdf", count, warnings, len(pages))


def filter_document(data: bytes, extension: str, redact: Callable[[str], dict]) -> DocumentResult:
    """Return a fresh filtered artifact, or fail the whole request without output."""
    if not isinstance(data, bytes) or not data:
        raise DocumentError(_EMPTY)
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(_TOO_LARGE, 413)
    if not isinstance(extension, str):
        raise DocumentError(_UNSUPPORTED, 422)
    extension = extension.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(_UNSUPPORTED, 422)
    try:
        if extension in {"xlsx", "docx"}:
            _zip_preflight(data, extension)
        elif extension == "pdf" and not data.startswith(b"%PDF-"):
            raise DocumentError(_MALFORMED)
        result = {"csv": _csv, "xlsx": _xlsx, "docx": _docx, "pdf": _pdf}[extension](data, redact)
        if len(result.content) > MAX_FILE_BYTES:
            raise DocumentError(_TOO_LARGE, 413)
        return result
    except DocumentError:
        raise
    except Exception:
        # Parser errors must never expose filenames, contents, OCR stderr or values.
        raise DocumentError(_MALFORMED) from None
