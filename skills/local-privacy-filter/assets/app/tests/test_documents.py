"""In-memory synthetic document tests; no records, network or model weights."""

import csv
import io
import logging
import re
import sys
import unittest
import zipfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from privacy_filter import documents
from privacy_filter.documents import DocumentError, filter_document


class FakeRedactor:
    def __init__(self):
        self.calls = []

    def __call__(self, text):
        self.calls.append(text)
        patterns = {
            "EMAIL_ADDRESS": r"[a-z]+@[a-z]+\.invalid",
            "PERSON": r"Alex Example|Sam Sample",
            "PHONE_NUMBER": r"0412 345 678",
            "DATE_OF_BIRTH": r"2000-01-02",
        }
        findings = [
            {"start": match.start(), "end": match.end(), "entity_type": label, "score": 0.9, "source": "synthetic"}
            for label, pattern in patterns.items() for match in re.finditer(pattern, text)
        ]
        for match in re.finditer(r"Name: (Taylor)\b", text):
            findings.append({"start": match.start(1), "end": match.end(1), "entity_type": "PERSON"})
        findings.sort(key=lambda item: item["start"])
        masked = text
        for item in reversed(findings):
            masked = masked[:item["start"]] + "[" + item["entity_type"] + "]" + masked[item["end"]:]
        return {"redacted_text": masked, "findings": findings, "entity_count": len(findings), "warnings": ["not trusted"]}


def docx_bytes(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def workbook_bytes(workbook):
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(text="Alex Example\nalex@example.invalid", image=False, blank=False):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image

    output = io.BytesIO()
    document = canvas.Canvas(output, invariant=1)
    document.setAuthor("Metadata Alex Example")
    document.setTitle("alex@example.invalid")
    if not blank:
        y = 700
        for line in text.splitlines():
            document.drawString(40, y, line)
            y -= 20
    if image:
        document.drawImage(ImageReader(Image.new("RGB", (10, 10), "white")), 40, 600, width=80, height=80)
    document.showPage()
    document.save()
    return output.getvalue()


def zip_xml_replace(data, name, replacement):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            target.writestr(item.filename, replacement if item.filename == name else original.read(item))
    return output.getvalue()


class DocumentFilteringTests(unittest.TestCase):
    def setUp(self):
        self.redact = FakeRedactor()

    def test_csv_preserves_values_structure_and_repeated_offsets(self):
        source = b'Name,Email,Notes\nAlex Example,alex@example.invalid,"ordinary, prose"\nSam Sample,alex@example.invalid,fine\n'
        original = bytes(source)
        result = filter_document(source, ".CSV", self.redact)
        rows = list(csv.reader(io.StringIO(result.content.decode())))
        self.assertEqual(rows, [["Name", "Email", "Notes"], ["[PERSON]", "[EMAIL_ADDRESS]", "ordinary, prose"], ["[PERSON]", "[EMAIL_ADDRESS]", "fine"]])
        self.assertEqual(result.entity_count, 4)
        self.assertEqual(result.units, 3)
        self.assertEqual(source, original)
        self.assertEqual(len(self.redact.calls), 1)
        self.assertNotIn("not trusted", result.warnings)

    def test_csv_header_context_masks_only_value(self):
        result = filter_document(b"Name,Description\nTaylor,plain text\n", "csv", self.redact)
        self.assertEqual(result.content, b"Name,Description\n[PERSON],plain text\n")

    def test_header_sensitive_value_filtered_but_context_not_copied(self):
        source = b"alex@example.invalid,Notes\nordinary,fine\n"
        result = filter_document(source, "csv", self.redact)
        self.assertEqual(result.entity_count, 1)
        self.assertEqual(result.content, b"[EMAIL_ADDRESS],Notes\nordinary,fine\n")

    def test_csv_utf16_delimiter_and_formula_neutralization(self):
        source = "Name;Formula;Notes\nAlex Example;=SUM(1);@example\n".encode("utf-16")
        result = filter_document(source, "csv", self.redact)
        rows = list(csv.reader(io.StringIO(result.content.decode())))
        self.assertEqual(rows[1], ["[PERSON]", "'=SUM(1)", "'@example"])

    def test_csv_batches_many_rows(self):
        source = "Name,Notes\n" + "Alex Example,ordinary prose\n" * 100
        result = filter_document(source.encode(), "csv", self.redact)
        self.assertEqual(result.entity_count, 100)
        self.assertLess(len(self.redact.calls), 5)

    def test_xlsx_rebuilds_hidden_cells_and_strips_metadata_links_comments(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment

        original = Workbook()
        original.active.title = "Alex Example"
        original.properties.creator = "alex@example.invalid"
        sheet = original.active
        sheet.append(["Name", "Email", "Date of birth", "Notes"])
        sheet.append(["Taylor", "alex@example.invalid", date(2000, 1, 2), "plain prose"])
        sheet["D2"].hyperlink = "https://example.invalid/alex@example.invalid"
        sheet["D2"].comment = Comment("Alex Example private comment", "alex@example.invalid")
        sheet.append(["Alex Example", "=\"alex@example.invalid\"", "=SUM(1)", "=safe literal"])
        hidden = original.create_sheet("Sam Sample")
        hidden.sheet_state = "hidden"
        hidden.append(["Name", "Phone"])
        hidden.append(["Sam Sample", "0412 345 678"])
        result = filter_document(workbook_bytes(original), "xlsx", self.redact)
        fresh = load_workbook(io.BytesIO(result.content), data_only=False)
        self.assertEqual(fresh.sheetnames, ["Sheet1", "Sheet2"])
        self.assertEqual(fresh["Sheet1"].max_row, 3)
        self.assertEqual(fresh["Sheet1"].max_column, 4)
        self.assertEqual(fresh["Sheet1"]["A2"].value, "[PERSON]")
        self.assertEqual(fresh["Sheet1"]["B2"].value, "[EMAIL_ADDRESS]")
        self.assertEqual(fresh["Sheet1"]["C2"].value, "[DATE_OF_BIRTH]")
        self.assertEqual(fresh["Sheet1"]["D2"].value, "plain prose")
        self.assertEqual(fresh["Sheet1"]["B3"].value, "[FORMULA OMITTED]")
        self.assertEqual(fresh["Sheet1"]["C3"].data_type, "s")
        self.assertEqual(fresh["Sheet2"]["A2"].value, "[PERSON]")
        self.assertEqual(fresh["Sheet2"]["B2"].value, "[PHONE_NUMBER]")
        self.assertEqual(fresh["Sheet2"].sheet_state, "visible")
        self.assertIsNone(fresh["Sheet1"]["D2"].comment)
        self.assertIsNone(fresh["Sheet1"]["D2"].hyperlink)
        self.assertIsNone(fresh.properties.creator)
        self.assertEqual(result.units, 5)
        self.assertEqual(result.entity_count, 6)
        with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
            self.assertFalse(any("comments" in name or "vba" in name or "externalLinks" in name for name in archive.namelist()))
            self.assertNotIn(b"alex@example.invalid", b"".join(archive.read(name) for name in archive.namelist()))

    def test_xlsx_string_formula_interpretation_is_disabled(self):
        from openpyxl import Workbook, load_workbook

        original = Workbook()
        original.active["A1"] = "Notes"
        original.active["A2"] = "=plain literal"
        original.active["A2"].data_type = "s"
        result = filter_document(workbook_bytes(original), "xlsx", self.redact)
        fresh = load_workbook(io.BytesIO(result.content))
        self.assertEqual(fresh.active["A2"].value, "=plain literal")
        self.assertEqual(fresh.active["A2"].data_type, "s")

    def test_docx_body_table_context_and_clean_package(self):
        from docx import Document

        original = Document()
        original.core_properties.author = "Alex Example"
        original.sections[0].header.paragraphs[0].text = "alex@example.invalid"
        original.sections[0].footer.paragraphs[0].text = "0412 345 678"
        original.add_paragraph("Ordinary introduction. Alex Example has an appointment.")
        table = original.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Phone"
        table.cell(1, 0).text = "Taylor"
        table.cell(1, 1).text = "0412 345 678"
        original.add_paragraph("Tail email: alex@example.invalid")
        result = filter_document(docx_bytes(original), "docx", self.redact)
        fresh = Document(io.BytesIO(result.content))
        self.assertEqual(fresh.paragraphs[0].text, "Ordinary introduction. [PERSON] has an appointment.")
        self.assertEqual(fresh.paragraphs[-1].text, "Tail email: [EMAIL_ADDRESS]")
        self.assertEqual(fresh.tables[0].cell(1, 0).text, "[PERSON]")
        self.assertEqual(fresh.tables[0].cell(1, 1).text, "[PHONE_NUMBER]")
        self.assertEqual(fresh.core_properties.author, "")
        self.assertEqual(fresh.sections[0].header.paragraphs[0].text, "")
        self.assertEqual(result.entity_count, 4)
        with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
            self.assertFalse(any(name.startswith("word/header") or name.startswith("word/footer") for name in archive.namelist()))
            self.assertNotIn(b"Alex Example", b"".join(archive.read(name) for name in archive.namelist()))

    def test_docx_cross_paragraph_name_context_is_retained(self):
        from docx import Document

        original = Document()
        original.add_paragraph("Name:")
        original.add_paragraph("Alex Example")
        original.add_paragraph("ordinary prose")
        result = filter_document(docx_bytes(original), "docx", self.redact)
        self.assertIn("Name:\nAlex Example\nordinary prose", self.redact.calls)
        self.assertEqual(Document(io.BytesIO(result.content)).paragraphs[1].text, "[PERSON]")

    def test_docx_insertions_flattened_deletions_and_text_boxes_omitted(self):
        from docx import Document
        from docx.oxml import OxmlElement

        original = Document()
        paragraph = original.add_paragraph("Ordinary ")
        insertion = OxmlElement("w:ins")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Alex Example"
        run.append(text)
        insertion.append(run)
        paragraph._p.append(insertion)
        deletion = OxmlElement("w:del")
        deleted = OxmlElement("w:delText")
        deleted.text = "private deleted secret"
        deletion.append(deleted)
        paragraph._p.append(deletion)
        result = filter_document(docx_bytes(original), "docx", self.redact)
        fresh = Document(io.BytesIO(result.content))
        self.assertEqual(fresh.paragraphs[0].text, "Ordinary [PERSON]")
        with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
            body = archive.read("word/document.xml")
            self.assertNotIn(b"private deleted secret", body)
            self.assertNotIn(b"<w:ins", body)
            self.assertNotIn(b"<w:del", body)

    def test_pdf_rebuilds_redacted_text_and_drops_original_properties(self):
        from pypdf import PdfReader

        result = filter_document(pdf_bytes(), "pdf", self.redact)
        fresh = PdfReader(io.BytesIO(result.content))
        text = "".join(page.extract_text() for page in fresh.pages)
        self.assertIn("[PERSON]", text)
        self.assertIn("[EMAIL_ADDRESS]", text)
        self.assertNotIn("Alex Example", text)
        self.assertNotIn("alex@example.invalid", str(fresh.metadata))
        self.assertEqual(fresh.metadata.get("/Author"), "")
        self.assertEqual(result.units, 1)
        self.assertEqual(result.entity_count, 2)

    def test_mixed_pdf_includes_ocr_and_readable_text_then_filters_both(self):
        from pypdf import PdfReader

        with patch.object(documents, "_ocr_page", return_value="Sam Sample\n0412 345 678") as ocr:
            result = filter_document(pdf_bytes(image=True), "pdf", self.redact)
        self.assertEqual(ocr.call_count, 1)
        text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(result.content)).pages)
        self.assertEqual(result.entity_count, 4)
        self.assertIn("[EMAIL_ADDRESS]", text)
        self.assertIn("[PHONE_NUMBER]", text)
        self.assertTrue(any("OCR" in warning for warning in result.warnings))

    def test_truly_blank_pdf_page_is_preserved_without_ocr(self):
        from pypdf import PdfWriter, PdfReader

        source = PdfWriter()
        source.add_blank_page(width=600, height=800)
        output = io.BytesIO()
        source.write(output)
        with patch.object(documents, "_ocr_page", side_effect=AssertionError("OCR should not run")):
            result = filter_document(output.getvalue(), "pdf", self.redact)
        self.assertEqual(len(PdfReader(io.BytesIO(result.content)).pages), 1)
        self.assertEqual(result.entity_count, 0)

    def test_blank_pdf_with_font_setup_is_preserved_without_ocr(self):
        with patch.object(documents, "_ocr_page", side_effect=AssertionError("OCR should not run")):
            result = filter_document(pdf_bytes(blank=True), "pdf", self.redact)
        self.assertEqual(result.entity_count, 0)
        self.assertEqual(result.units, 1)

    def test_malformed_pdf_font_cannot_write_identifiers_to_parser_logs(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, TextStringObject

        source = PdfReader(io.BytesIO(pdf_bytes()))
        font = source.pages[0]["/Resources"]["/Font"]["/F1"].get_object()
        font[NameObject("/Encoding")] = DictionaryObject({
            NameObject("/BaseEncoding"): NameObject("/WinAnsiEncoding"),
            NameObject("/Differences"): TextStringObject("alex@example.invalid"),
        })
        writer = PdfWriter()
        writer.add_page(source.pages[0])
        output = io.BytesIO()
        writer.write(output)
        # The real parser otherwise logs the raw malformed Differences value.
        logging.getLogger("pypdf._cmap").disabled = False
        with self.assertNoLogs("pypdf", level="WARNING"):
            result = filter_document(output.getvalue(), "pdf", self.redact)
        self.assertEqual(result.entity_count, 2)

    def test_pdf_decompression_limit_blocks_large_content_stream(self):
        with patch.object(documents, "MAX_PDF_STREAM_BYTES", 20), self.assertRaises(DocumentError) as caught:
            filter_document(pdf_bytes(), "pdf", self.redact)
        self.assertIn(str(caught.exception), documents.SAFE_ERROR_MESSAGES)

    def test_image_pdf_missing_ocr_and_ocr_errors_fail_whole_request(self):
        for message in (documents._OCR_MISSING, documents._OCR_FAILED):
            with self.subTest(message=message), patch.object(documents, "_ocr_page", side_effect=DocumentError(message, 422)):
                with self.assertRaises(DocumentError) as caught:
                    filter_document(pdf_bytes(image=True), "pdf", self.redact)
                self.assertEqual(caught.exception.status, 422)
                self.assertIn(str(caught.exception), documents.SAFE_ERROR_MESSAGES)

    def test_ocr_pipes_bytes_with_bounded_timeouts_and_no_content_paths(self):
        png = b"\x89PNG\r\n\x1a\n" + b"synthetic"
        with patch.object(documents, "_ocr_binaries", return_value=("/local/pdftoppm", "/local/tesseract")), patch.object(documents.subprocess, "run", side_effect=[SimpleNamespace(stdout=png), SimpleNamespace(stdout=b"Alex Example")]) as runner:
            result = documents._ocr_page(b"%PDF-synthetic", 3, documents.time.monotonic() + 120)
        self.assertEqual(result, "Alex Example")
        first, second = runner.call_args_list
        self.assertEqual(first.args[0], ["/local/pdftoppm", "-f", "3", "-l", "3", "-scale-to", "2000", "-singlefile", "-png", "-"])
        self.assertEqual(first.kwargs["input"], b"%PDF-synthetic")
        self.assertEqual(second.args[0], ["/local/tesseract", "stdin", "stdout", "-l", "eng", "--psm", "3"])
        self.assertEqual(second.kwargs["input"], png)
        self.assertLessEqual(first.kwargs["timeout"], 30)
        self.assertNotIn("shell", first.kwargs)
        self.assertEqual(first.kwargs["stderr"], documents.subprocess.DEVNULL)

    def test_blank_ocr_never_returns_partial_file(self):
        png = b"\x89PNG\r\n\x1a\n"
        with patch.object(documents, "_ocr_binaries", return_value=("pdftoppm", "tesseract")), patch.object(documents.subprocess, "run", side_effect=[SimpleNamespace(stdout=png), SimpleNamespace(stdout=b"  \n")]):
            with self.assertRaises(DocumentError) as caught:
                documents._ocr_page(b"%PDF", 1, documents.time.monotonic() + 120)
        self.assertEqual(str(caught.exception), documents._OCR_FAILED)

    def test_encrypted_pdf_rejected(self):
        from pypdf import PdfReader, PdfWriter

        writer = PdfWriter()
        writer.add_page(PdfReader(io.BytesIO(pdf_bytes())).pages[0])
        writer.encrypt("synthetic-password")
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaises(DocumentError) as caught:
            filter_document(data.getvalue(), "pdf", self.redact)
        self.assertEqual(str(caught.exception), documents._ENCRYPTED)

    def test_rejects_unsupported_empty_malformed_and_upload_limit(self):
        for data, extension in ((b"legacy", "doc"), (b"legacy", "xls"), (b"macro", "xlsm"), (b"", "csv"), (b"broken", "pdf"), (b"not a zip", "docx")):
            with self.subTest(extension=extension), self.assertRaises(DocumentError) as caught:
                filter_document(data, extension, self.redact)
            self.assertIn(str(caught.exception), documents.SAFE_ERROR_MESSAGES)
        with patch.object(documents, "MAX_FILE_BYTES", 4), self.assertRaises(DocumentError) as caught:
            filter_document(b"12345", "csv", self.redact)
        self.assertEqual(caught.exception.status, 413)

    def test_filter_failures_and_invalid_spans_fail_without_input_in_error(self):
        source = b"Alex Example\n"
        bad_results = ({"redacted_text": "fine"}, {"findings": [{"start": 0, "end": 99, "entity_type": "PERSON"}]}, {"findings": [{"start": True, "end": 4, "entity_type": "PERSON"}]}, {"findings": [{"start": 0, "end": 4, "entity_type": "Alex Example"}]})
        for bad in bad_results:
            with self.subTest(bad=bad), self.assertRaises(DocumentError) as caught:
                filter_document(source, "csv", lambda _: bad)
            self.assertEqual(str(caught.exception), documents._FILTER_FAILED)
        def broken(_):
            raise RuntimeError("Alex Example private error")
        with self.assertRaises(DocumentError) as caught:
            filter_document(source, "csv", broken)
        self.assertNotIn("Alex Example", str(caught.exception))
        self.assertEqual(caught.exception.status, 503)

    def test_overlap_union_prevents_partial_value_leakage(self):
        text = "Alex Example"
        def spans(_):
            return {"findings": [{"start": 0, "end": 7, "entity_type": "PERSON"}, {"start": 5, "end": len(text), "entity_type": "LOCATION"}]}
        result = filter_document(text.encode(), "csv", spans)
        self.assertEqual(result.content, b"[SENSITIVE_DATA]\n")
        self.assertEqual(result.entity_count, 1)

    def test_document_limits_include_context_and_rows(self):
        with patch.object(documents, "MAX_DOCUMENT_CHARS", 20), self.assertRaises(DocumentError) as caught:
            filter_document(b"LongHeader\n1234\n1234\n", "csv", self.redact)
        self.assertEqual(caught.exception.status, 413)
        with patch.object(documents, "MAX_ROWS", 2), self.assertRaises(DocumentError):
            filter_document(b"Name\nAlex Example\nSam Sample\n", "csv", self.redact)

    def test_zip_bomb_macro_xml_entity_and_traversal_rejected(self):
        def zip_data(entries):
            output = io.BytesIO()
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                for name, value in entries:
                    archive.writestr(name, value)
            return output.getvalue()
        cases = [
            [("../private.xml", "<r/>"), ("word/document.xml", "<r/>"), ("[Content_Types].xml", "<r/>")],
            [("word/vbaProject.bin", "macro"), ("word/document.xml", "<r/>"), ("[Content_Types].xml", "<r/>")],
            [("word/document.xml", '<!DOCTYPE r [<!ENTITY x "Alex Example">]><r>&x;</r>'), ("[Content_Types].xml", "<r/>")],
            [("word/document.xml", "<r>" + "a" * 100_000 + "</r>"), ("[Content_Types].xml", "<r/>")],
        ]
        for entries in cases:
            with self.subTest(entries=entries[0][0]), self.assertRaises(DocumentError):
                filter_document(zip_data(entries), "docx", self.redact)

    def test_xlsx_huge_sparse_coordinate_rejected_before_parser(self):
        from openpyxl import Workbook

        book = Workbook()
        book.active["A1"] = "ordinary"
        source = workbook_bytes(book)
        with zipfile.ZipFile(io.BytesIO(source)) as archive:
            worksheet = archive.read("xl/worksheets/sheet1.xml").replace(b'r="A1"', b'r="A1000000"')
        source = zip_xml_replace(source, "xl/worksheets/sheet1.xml", worksheet)
        with self.assertRaises(DocumentError) as caught:
            filter_document(source, "xlsx", self.redact)
        self.assertEqual(caught.exception.status, 413)


if __name__ == "__main__":
    unittest.main()
